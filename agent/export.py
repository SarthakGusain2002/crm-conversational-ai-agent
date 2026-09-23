"""Export CRM search results to Excel, Word, or PDF.

Each function takes the same shape of input (a list of flat dict records)
and returns raw bytes ready to hand to a download button — no files are
written to disk.
"""
from __future__ import annotations

from io import BytesIO

import pandas as pd
from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


def _as_dataframe(records: list[dict]) -> pd.DataFrame:
    return pd.DataFrame(records)


def to_excel_bytes(records: list[dict]) -> bytes:
    df = _as_dataframe(records)
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Results")
    return buffer.getvalue()


def to_word_bytes(records: list[dict], title: str = "CRM Search Results") -> bytes:
    document = Document()
    heading = document.add_heading(title, level=1)
    heading.style.font.size = Pt(18)

    if not records:
        document.add_paragraph("No matching records found.")
    else:
        columns = list(records[0].keys())
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Light Grid Accent 1"
        header_cells = table.rows[0].cells
        for i, col in enumerate(columns):
            header_cells[i].text = str(col)
        for record in records:
            row_cells = table.add_row().cells
            for i, col in enumerate(columns):
                row_cells[i].text = str(record.get(col, ""))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def to_pdf_bytes(records: list[dict], title: str = "CRM Search Results") -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    if not records:
        story.append(Paragraph("No matching records found.", styles["Normal"]))
    else:
        columns = list(records[0].keys())
        data = [columns] + [[str(record.get(col, "")) for col in columns] for record in records]
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3a6b")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f6fb")]),
                ]
            )
        )
        story.append(table)

    doc.build(story)
    return buffer.getvalue()
